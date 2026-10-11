# -*- coding: utf-8 -*-
"""真实验收：用最脏的数据（群 439880516 / 私聊 3048268426）导出并校验。

覆盖：
  1) 群 439880516 → docx 完整导出（含 \x14 脏数据）
  2) 群 439880516 → md  完整导出，打印前 20 行 + \x14 所在日期段落
  3) 私聊 3048268426 → docx 导出
  4) 取消测试：不留下半截文件
  5) docx 降级测试：python-docx 写入中途抛异常 → 自动退化 OOXML

结果写入 exports/acceptance.txt（同名 .done 为完成标记）。
"""
from __future__ import annotations

import os
import sys
import threading
import time
import zipfile

STUDIO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if STUDIO not in sys.path:
    sys.path.insert(0, STUDIO)

from core.config import Config                 # noqa: E402
from core.exporter import (                    # noqa: E402
    _DocxWriter, default_filename, export_chat,
)
from core.memory import MemoryStore            # noqa: E402

OUT_DIR = os.path.join(STUDIO, "exports")
LOG: list[str] = []


def emit(s: str = "") -> None:
    print(s, flush=True)
    LOG.append(s)


def chat_name(store, ct, cid) -> str:
    for c in store.chats():
        if c["chat_type"] == ct and c["chat_id"] == cid:
            return c["name"]
    return cid


def verify_docx(path: str, expect_count: int) -> None:
    from docx import Document
    d = Document(path)
    paras = d.paragraphs
    heads = [p for p in paras if p.style.name.startswith("Heading")]
    n_h1 = sum(1 for p in heads if p.style.name == "Heading 1")
    n_h2 = sum(1 for p in heads if p.style.name == "Heading 2")
    # 段落总数 = 1(h1) + 5(元信息) + n_h2(日期标题) + 正文条数
    expect_paras = 1 + 5 + n_h2 + expect_count
    emit(f"  [docx] 段落总数={len(paras):,}  标题={len(heads)}（H1={n_h1} H2={n_h2}）")
    emit(f"  [docx] 期望段落数=1+5+{n_h2}+{expect_count:,}={expect_paras:,}  "
         f"实际={len(paras):,}  {'✔一致' if len(paras) == expect_paras else '✘不一致'}")
    emit(f"  [docx] 前 3 段：{[p.text[:60] for p in paras[:3]]}")
    # 确认脏字符已被清洗、文档内不再有裸露控制符
    joined = "\n".join(p.text for p in paras)
    bad = [c for c in joined if (ord(c) <= 0x1F and c not in "\t\n\r") or 0xD800 <= ord(c) <= 0xDFFF
           or ord(c) in (0xFFFE, 0xFFFF)]
    emit(f"  [docx] 文档内残留非法字符数={len(bad)}  {'✔干净' if not bad else '✘仍有残留'}")
    if "␔/泪奔" in joined:
        emit("  [docx] 命中已清洗的脏正文占位：……␔/泪奔 ✔")


def main() -> int:
    os.makedirs(OUT_DIR, exist_ok=True)
    cfg = Config.load()
    store = MemoryStore(cfg.memory_dir)
    emit(f"记忆库：{cfg.memory_dir}")
    emit(f"开始：{time.strftime('%Y-%m-%d %H:%M:%S')}")

    # =============================================================== 1) 群 docx
    CT, CID = "group", "439880516"
    name = chat_name(store, CT, CID)
    cnt = store.count_chat_messages(CT, CID)
    emit("")
    emit("=" * 70)
    emit(f"[1] 群 {CID}（{name}）→ docx   共 {cnt:,} 条")
    docx_path = os.path.join(OUT_DIR, "accept_" + default_filename(name, CT, "docx"))
    if os.path.exists(docx_path):
        os.remove(docx_path)
    t0 = time.time()
    _last = {"n": 0}
    def _prog(d, t, e):
        if d - _last["n"] >= 20000:
            _last["n"] = d
            emit(f"    [prog] {d:,}/{t:,} {e:.1f}s")
    res = export_chat(store, {"chat_type": CT, "chat_id": CID, "chat_name": name, "fmt": "docx",
                              "out_path": docx_path, "by_day": True, "show_qq": True},
                      on_log=lambda s: emit("    [log] " + s),
                      on_progress=_prog)
    dt = time.time() - t0
    emit(f"  => ok={res['ok']} count={res['count']:,} cleaned={res['cleaned']:,} "
         f"bytes={res['bytes']:,} elapsed={res['elapsed']}s degraded={res['degraded']} "
         f"err={res['error']}")
    emit(f"  => 绝对路径：{os.path.abspath(docx_path)}")
    emit(f"  => 文件大小：{os.path.getsize(docx_path):,} 字节  ({os.path.getsize(docx_path)/1024/1024:.2f} MB)")
    emit(f"  => 实测墙钟耗时：{dt:.1f}s")
    emit(f"  => 半成品 .part 是否残留：{os.path.exists(docx_path + '.part')}")
    verify_docx(docx_path, res["count"])

    # =============================================================== 2) 群 md
    emit("")
    emit("=" * 70)
    emit(f"[2] 群 {CID}（{name}）→ md   共 {cnt:,} 条")
    md_path = os.path.join(OUT_DIR, "accept_" + default_filename(name, CT, "md"))
    if os.path.exists(md_path):
        os.remove(md_path)
    res2 = export_chat(store, {"chat_type": CT, "chat_id": CID, "chat_name": name, "fmt": "md",
                               "out_path": md_path, "by_day": True, "show_qq": True},
                       on_log=lambda s: None)
    emit(f"  => ok={res2['ok']} count={res2['count']:,} cleaned={res2['cleaned']:,} "
         f"bytes={res2['bytes']:,} elapsed={res2['elapsed']}s err={res2['error']}")
    emit(f"  => 绝对路径：{os.path.abspath(md_path)}")
    emit("  [md] 前 20 行：")
    with open(md_path, "r", encoding="utf-8") as f:
        for i, line in enumerate(f):
            if i >= 20:
                break
            emit("     |" + line.rstrip("\n"))
    # 定位脏正文所在日期段落
    emit("  [md] 含占位 ␔ 的日期段落上下文（\"……␔/泪奔\"）：")
    with open(md_path, "r", encoding="utf-8") as f:
        text = f.read()
    idx = text.find("␔/泪奔")
    emit(f"     未找到" if idx < 0 else "     找到")
    if idx >= 0:
        seg = text[max(0, idx - 400):idx + 200]
        # 找到该行所属的最近日期标题
        head = text.rfind("\n## ", 0, idx)
        day_title = text[head + 1: text.find("\n", head + 1)] if head >= 0 else "?"
        emit(f"     所属日期标题：{day_title}")
        emit("     上下文：")
        for ln in seg.splitlines():
            emit("       > " + ln)
    # 确认 md 里不含裸控制字符
    bad = [c for c in text if (ord(c) <= 0x1F and c not in "\t\n\r")
           or 0xD800 <= ord(c) <= 0xDFFF or ord(c) in (0xFFFE, 0xFFFF)]
    emit(f"  [md] 全文残留非法控制字符数={len(bad)}  {'✔干净' if not bad else '✘仍有残留'}")

    # =============================================================== 3) 私聊 docx
    emit("")
    emit("=" * 70)
    CT2, CID2 = "c2c", "3048268426"
    name2 = chat_name(store, CT2, CID2)
    cnt2 = store.count_chat_messages(CT2, CID2)
    emit(f"[3] 私聊 {CID2}（{name2}）→ docx   共 {cnt2:,} 条")
    docx2 = os.path.join(OUT_DIR, "accept_" + default_filename(name2, CT2, "docx"))
    if os.path.exists(docx2):
        os.remove(docx2)
    res3 = export_chat(store, {"chat_type": CT2, "chat_id": CID2, "chat_name": name2, "fmt": "docx",
                               "out_path": docx2, "by_day": True, "show_qq": True},
                       on_log=lambda s: None)
    emit(f"  => ok={res3['ok']} count={res3['count']:,} cleaned={res3['cleaned']:,} "
         f"bytes={res3['bytes']:,} elapsed={res3['elapsed']}s err={res3['error']}")
    emit(f"  => 绝对路径：{os.path.abspath(docx2)}")
    emit(f"  => 文件大小：{os.path.getsize(docx2):,} 字节")
    verify_docx(docx2, res3["count"])

    # =============================================================== 4) 取消测试
    emit("")
    emit("=" * 70)
    emit("[4] 取消测试：导出群 439880516 → txt，0.4s 后取消")
    cancel_path = os.path.join(OUT_DIR, "accept_cancel_439880516.txt")
    for p in (cancel_path, cancel_path + ".part"):
        if os.path.exists(p):
            os.remove(p)
    ev = threading.Event()
    threading.Timer(0.4, ev.set).start()
    res4 = export_chat(store, {"chat_type": CT, "chat_id": CID, "chat_name": name, "fmt": "txt",
                               "out_path": cancel_path, "by_day": True},
                       on_log=lambda s: emit("    [log] " + s), should_cancel=ev.is_set)
    emit(f"  => canceled={res4['canceled']} 已写={res4['count']:,} ok={res4['ok']}")
    emit(f"  => 成品文件存在={os.path.exists(cancel_path)}  .part 残留={os.path.exists(cancel_path + '.part')}"
         f"  {'✔未留半成品' if not os.path.exists(cancel_path) and not os.path.exists(cancel_path + '.part') else '✘有残留'}")

    # =============================================================== 5) docx 降级
    emit("")
    emit("=" * 70)
    emit("[5] docx 降级测试：python-docx 写入中途抛异常 → 自动退化 OOXML")
    fb_path = os.path.join(OUT_DIR, "accept_fallback.docx")
    if os.path.exists(fb_path):
        os.remove(fb_path)
    orig = _DocxWriter._emit_docx
    state = {"n": 0}
    def boom(self, text, level):
        state["n"] += 1
        if state["n"] == 50:          # 第 50 段时模拟 python-docx 崩掉
            raise ValueError("All strings must be XML compatible: 模拟故障")
        return orig(self, text, level)
    _DocxWriter._emit_docx = boom
    try:
        res5 = export_chat(store, {"chat_type": CT2, "chat_id": CID2, "chat_name": name2,
                                   "fmt": "docx", "out_path": fb_path, "by_day": True,
                                   "show_qq": True},
                           on_log=lambda s: emit("    [log] " + s))
    finally:
        _DocxWriter._emit_docx = orig
    emit(f"  => ok={res5['ok']} count={res5['count']:,} degraded={res5['degraded']} "
         f"bytes={res5['bytes']:,} err={res5['error']}")
    emit(f"  => 绝对路径：{os.path.abspath(fb_path)}")
    # 用 python-docx 打开降级产物，验证互操作性
    try:
        from docx import Document
        dd = Document(fb_path)
        emit(f"  => 降级文件可被 python-docx 打开：段落数={len(dd.paragraphs):,}")
    except Exception as e:  # noqa: BLE001
        emit(f"  => 降级文件打开失败：{type(e).__name__}: {e}")
    with zipfile.ZipFile(fb_path) as z:
        import xml.dom.minidom as minidom
        minidom.parseString(z.read("word/document.xml"))
    emit("  => word/document.xml XML 合法性 ✔")

    emit("")
    emit(f"完成：{time.strftime('%Y-%m-%d %H:%M:%S')}")

    with open(os.path.join(OUT_DIR, "acceptance.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(LOG) + "\n")
    with open(os.path.join(OUT_DIR, "acceptance.txt.done"), "w", encoding="utf-8") as f:
        f.write("done\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
